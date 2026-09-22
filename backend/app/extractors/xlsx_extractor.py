from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.extractors.base import BaseExtractor, ExtractedElement, ExtractionMetadata, ExtractionResult

MAX_ROWS_RENDERED = 300


class XLSXExtractor(BaseExtractor):
    """One `sheet` element per worksheet, carrying structured rows/headers
    in `extra` so SpreadsheetChunker can chunk by logical table rather than
    by character count (spec section 19)."""

    name = "XLSXExtractor"
    version = "1.0"

    def can_handle(self, parser_type: str) -> bool:
        return parser_type == "xlsx"

    def extract(self, content: bytes, filename: str) -> ExtractionResult:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
        elements: list[ExtractedElement] = []
        full_text_parts: list[str] = []

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows: list[list[str]] = []
                for row in ws.iter_rows(values_only=True):
                    if row is None or all(v is None for v in row):
                        continue
                    rows.append(["" if v is None else str(v) for v in row])
                if not rows:
                    continue
                headers, data_rows = rows[0], rows[1:]
                rendered_lines = [" | ".join(headers)] + [" | ".join(r) for r in data_rows[:MAX_ROWS_RENDERED]]
                rendered = "\n".join(rendered_lines)
                elements.append(
                    ExtractedElement(
                        type="sheet",
                        text=rendered,
                        sheet=sheet_name,
                        extra={"headers": headers, "rows": data_rows, "row_count": len(data_rows)},
                    )
                )
                full_text_parts.append(rendered)

            metadata = ExtractionMetadata(sheet_names=wb.sheetnames, word_count=sum(len(t.split()) for t in full_text_parts))
        finally:
            wb.close()

        return ExtractionResult(
            elements=elements,
            metadata=metadata,
            full_text="\n\n".join(full_text_parts),
            extractor_name=self.name,
            extractor_version=self.version,
        )
