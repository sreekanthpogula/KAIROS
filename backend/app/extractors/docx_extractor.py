from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument

from app.extractors.base import BaseExtractor, ExtractedElement, ExtractionMetadata, ExtractionResult


class DOCXExtractor(BaseExtractor):
    name = "DOCXExtractor"
    version = "1.0"

    def can_handle(self, parser_type: str) -> bool:
        return parser_type == "docx"

    def extract(self, content: bytes, filename: str) -> ExtractionResult:
        doc = DocxDocument(BytesIO(content))
        elements: list[ExtractedElement] = []
        full_text_parts: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = (para.style.name or "").lower() if para.style else ""
            if style_name.startswith("heading") or style_name == "title":
                digits = "".join(c for c in style_name if c.isdigit())
                level = int(digits) if digits else 1
                elements.append(ExtractedElement(type="heading", text=text, heading_level=level))
            else:
                elements.append(ExtractedElement(type="paragraph", text=text))
            full_text_parts.append(text)

        for t_index, table in enumerate(doc.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            rows = [r for r in rows if any(c for c in r)]
            if not rows:
                continue
            rendered = "\n".join(" | ".join(r) for r in rows)
            elements.append(ExtractedElement(type="table", text=rendered, extra={"rows": rows, "table_index": t_index}))
            full_text_parts.append(rendered)

        props = doc.core_properties
        metadata = ExtractionMetadata(
            title=props.title or None,
            author=props.author or None,
            created=props.created.isoformat() if props.created else None,
            word_count=sum(len(t.split()) for t in full_text_parts),
        )

        return ExtractionResult(
            elements=elements,
            metadata=metadata,
            full_text="\n\n".join(full_text_parts),
            extractor_name=self.name,
            extractor_version=self.version,
        )
