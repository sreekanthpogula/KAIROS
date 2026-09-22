from __future__ import annotations

from io import BytesIO

from pptx import Presentation

from app.extractors.base import BaseExtractor, ExtractedElement, ExtractionMetadata, ExtractionResult


class PPTXExtractor(BaseExtractor):
    """One `heading` (slide title) plus one `slide` element (body text) per
    slide — the natural chunk boundary for PresentationChunker."""

    name = "PPTXExtractor"
    version = "1.0"

    def can_handle(self, parser_type: str) -> bool:
        return parser_type == "pptx"

    def extract(self, content: bytes, filename: str) -> ExtractionResult:
        prs = Presentation(BytesIO(content))
        elements: list[ExtractedElement] = []
        full_text_parts: list[str] = []

        for idx, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            title_text: str | None = None
            body_texts: list[str] = []

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if not text.strip():
                    continue
                is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
                if is_title:
                    title_text = text
                else:
                    body_texts.append(text)

            if title_text:
                elements.append(ExtractedElement(type="heading", text=title_text, slide=idx, heading_level=1))
                full_text_parts.append(title_text)

            body_combined = "\n".join(body_texts).strip()
            if body_combined:
                elements.append(ExtractedElement(type="slide", text=body_combined, slide=idx))
                full_text_parts.append(body_combined)

        metadata = ExtractionMetadata(
            slide_count=len(prs.slides),
            word_count=sum(len(t.split()) for t in full_text_parts),
        )

        return ExtractionResult(
            elements=elements,
            metadata=metadata,
            full_text="\n\n".join(full_text_parts),
            extractor_name=self.name,
            extractor_version=self.version,
        )
