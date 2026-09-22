"""Generates the synthetic ECIP demo corpus into data/samples/.

Every document is entirely synthetic enterprise content authored for this
POC. No real patient data, no real PHI, no data pulled from any external
source — see docs/decisions.md for why (ADR on synthetic-only corpus) and
docs/evaluation.md for how the "1TB" Scale Simulator is instead calibrated
against publicly documented reference statistics rather than a real corpus.

Run: python scripts/seed_demo_data.py
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = REPO_ROOT / "data" / "samples"


# ---------------------------------------------------------------------------
# Format writers
# ---------------------------------------------------------------------------
def write_pdf(path: Path, title: str, sections: list[tuple[str, str]], table: tuple[list, list[list]] | None = None) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ECIPTitle", parent=styles["Title"], fontSize=20, spaceAfter=18)
    heading_style = ParagraphStyle("ECIPHeading", parent=styles["Heading1"], fontSize=16, spaceAfter=10, spaceBefore=14)
    body_style = ParagraphStyle("ECIPBody", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=8)

    flow = [Paragraph(title, title_style), Spacer(1, 12)]
    for heading, body in sections:
        flow.append(Paragraph(heading, heading_style))
        for para in body.split("\n\n"):
            flow.append(Paragraph(para.replace("\n", "<br/>"), body_style))
    if table:
        header, rows = table
        flow.append(Spacer(1, 8))
        t = Table([header] + rows, hAlign="LEFT")
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dddddd")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        flow.append(t)
    doc = SimpleDocTemplate(str(path), pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    doc.build(flow)


def write_docx(path: Path, title: str, sections: list[tuple[str, str]], table: tuple[list, list[list]] | None = None) -> None:
    doc = DocxDocument()
    doc.add_heading(title, level=0)
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        for para in body.split("\n\n"):
            doc.add_paragraph(para)
    if table:
        header, rows = table
        t = doc.add_table(rows=1, cols=len(header))
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = h
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
    doc.save(str(path))


def write_xlsx(path: Path, sheets: dict[str, tuple[list, list[list]]]) -> None:
    wb = Workbook()
    first = True
    for name, (headers, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        ws.append(headers)
        for row in rows:
            ws.append(row)
    wb.save(str(path))


def write_pptx(path: Path, slides: list[tuple[str, list[str]]]) -> None:
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]

    title, subtitle_bullets = slides[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "\n".join(subtitle_bullets)

    for slide_title, bullets in slides[1:]:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_title
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            p = tf.add_paragraph()
            p.text = b
    prs.save(str(path))


def write_html(
    path: Path,
    title: str,
    sections: list[tuple[str, int, str | list[str]]],
    table: tuple[list, list[list]] | None = None,
) -> None:
    parts = [f"<h1>{title}</h1>"]
    for heading, level, content in sections:
        parts.append(f"<h{level}>{heading}</h{level}>")
        if isinstance(content, list):
            parts.append("<ul>" + "".join(f"<li>{c}</li>" for c in content) + "</ul>")
        else:
            for para in content.split("\n\n"):
                parts.append(f"<p>{para}</p>")
    if table:
        header, rows = table
        parts.append("<table border='1'>")
        parts.append("<tr>" + "".join(f"<th>{h}</th>" for h in header) + "</tr>")
        for row in rows:
            parts.append("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>")
        parts.append("</table>")
    html = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{title}</title></head><body>{''.join(parts)}</body></html>"
    path.write_text(html, encoding="utf-8")


def write_text(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def write_csv(path: Path, headers: list[str], rows: list[list]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)


def write_eml(path: Path, sender: str, to: str, subject: str, date: str, body: str) -> None:
    content = (
        f"From: {sender}\r\nTo: {to}\r\nSubject: {subject}\r\nDate: {date}\r\n"
        f"MIME-Version: 1.0\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n{body}\r\n"
    )
    path.write_bytes(content.encode("utf-8"))


# ---------------------------------------------------------------------------
# Document content — 18 synthetic enterprise documents, 9 file formats
# ---------------------------------------------------------------------------
GOLDEN_LABELS: list[dict] = []


def _label(filename: str, ontology_id: str, domain: str, document_type: str, document_subtype: str) -> None:
    GOLDEN_LABELS.append(
        {
            "filename": filename,
            "ontology_id": ontology_id,
            "domain": domain,
            "document_type": document_type,
            "document_subtype": document_subtype,
        }
    )


def build_provider_agreement() -> None:
    fn = "provider_agreement_northvalley.pdf"
    sections = [
        ("Section 1: Definitions", "This Provider Agreement (\"Agreement\") is entered into by and between North Valley Hospital (\"Provider\") and Meridian Health Partners (\"Payer\"), effective January 1, 2026.\n\n\"Covered Services\" means those health care services reimbursable under this Agreement. \"Effective Date\" means January 1, 2026."),
        ("Section 2: Payment", "Payer shall reimburse Provider in accordance with the Fee Schedule attached as Appendix A. Payment shall be issued within forty-five (45) days of receipt of a clean claim. Reimbursement rate reviews occur annually."),
        ("Section 3: Eligibility", "Provider must maintain active credentialing status and network participation in good standing to remain eligible for reimbursement under this Agreement. Provider eligibility is reviewed quarterly by the Provider Management team."),
        ("Section 4: Termination", "Either party may terminate this Agreement upon ninety (90) days written notice to the other party. Termination for cause, including loss of licensure or a material breach of the Compliance provisions in Section 5, may occur upon thirty (30) days written notice. Provider termination requires written notice delivered to the Provider Network Management department."),
        ("Section 5: Compliance", "Provider shall comply with all applicable HIPAA privacy and security requirements and all state and federal health care regulations. Provider shall maintain PHI safeguards consistent with Provider's HIPAA Compliance Policy."),
        ("Section 6: Appendices", "Appendix A: Fee Schedule. Appendix B: Covered Service Codes. Appendix C: Credentialing Requirements."),
    ]
    write_pdf(OUT_DIR / fn, "Provider Agreement", sections)
    _label(fn, "legal.contracts.provider_agreement", "legal", "contract", "provider_agreement")


def build_employee_healthcare_policy() -> None:
    fn = "employee_healthcare_policy.docx"
    sections = [
        ("Purpose", "This Employee Healthcare Policy establishes the health benefits available to eligible employees of the organization and the enrollment procedures required to participate."),
        ("Eligibility", "Full-time employees become eligible for the employee healthcare policy on the first day of the month following sixty (60) days of continuous employment. Dependents may be added during open enrollment."),
        ("Enrollment", "Employees must complete enrollment through the HR benefits portal within thirty (30) days of becoming eligible. Failure to enroll within this window requires a qualifying life event to enroll outside the open enrollment period."),
        ("Plan Options", "The organization offers three plan tiers described in the table below."),
        ("Employee Responsibilities", "Employees are responsible for reviewing plan documents, reporting qualifying life events within thirty (30) days, and paying applicable premium contributions via payroll deduction."),
        ("Contact", "Questions regarding this employee healthcare policy should be directed to the Human Resources Benefits team."),
    ]
    table = (["Plan Name", "Monthly Premium", "Deductible"], [["Plan A - HMO", "$85", "$500"], ["Plan B - PPO", "$140", "$1,000"], ["Plan C - HDHP", "$60", "$2,500"]])
    write_docx(OUT_DIR / fn, "Employee Healthcare Policy", sections, table=table)
    _label(fn, "administrative.policies.employee_healthcare_policy", "administrative", "policy", "employee_healthcare_policy")


def build_patient_care_guideline() -> None:
    fn = "patient_care_guideline_diabetes.docx"
    sections = [
        ("Purpose", "This patient care guideline establishes the standard of care for outpatient management of adult patients with Type 2 diabetes."),
        ("Scope", "This clinical guideline applies to primary care providers and endocrinology staff across all outpatient clinics."),
        ("Clinical Recommendations", "Providers should target an HbA1c below 7% for most adults, individualized based on comorbidities. Metformin remains the first-line pharmacologic therapy absent contraindication. Evidence-based practice supports annual retinal and nephropathy screening."),
        ("Monitoring", "Patients should have HbA1c reviewed every three months until stable, then every six months. Blood pressure and lipid panels should be reviewed at each visit per the monitoring schedule below."),
        ("References", "American Diabetes Association Standards of Care; internal Clinical Guidelines Committee review, updated 2026."),
    ]
    table = (["Test", "Frequency"], [["HbA1c", "Every 3-6 months"], ["Blood Pressure", "Every visit"], ["Lipid Panel", "Annually"], ["Retinal Exam", "Annually"]])
    write_docx(OUT_DIR / fn, "Patient Care Guideline: Type 2 Diabetes Management", sections, table=table)
    _label(fn, "clinical.clinical_guidelines.patient_care_guideline", "clinical", "clinical_guideline", "patient_care_guideline")


def build_insurance_claim_report() -> None:
    fn = "insurance_claim_report_q3.xlsx"
    headers = ["Claim ID", "Provider", "Payer", "CPT Code", "Billed Amount", "Paid Amount", "Status", "Date"]
    rows = [
        ["CLM-100234", "North Valley Hospital", "Meridian Health Partners", "99213", 185.00, 148.00, "Paid", "2026-07-03"],
        ["CLM-100235", "North Valley Hospital", "Meridian Health Partners", "99214", 245.00, 196.00, "Paid", "2026-07-05"],
        ["CLM-100236", "North Valley Hospital", "Summit Care Network", "80053", 95.00, 0.00, "Denied", "2026-07-08"],
        ["CLM-100237", "North Valley Hospital", "Meridian Health Partners", "99396", 210.00, 168.00, "Paid", "2026-07-12"],
        ["CLM-100238", "North Valley Hospital", "Summit Care Network", "36415", 25.00, 20.00, "Paid", "2026-07-14"],
    ]
    write_xlsx(OUT_DIR / fn, {"Claims": (headers, rows)})
    _label(fn, "financial.claims.insurance_claim_report", "financial", "claim_report", "insurance_claim_report")


def build_hospital_invoice() -> None:
    fn = "hospital_invoice_2456.pdf"
    sections = [
        ("Invoice", "Invoice Number: INV-2456\nBill To: Meridian Health Partners, Accounts Payable\nInvoice Date: August 1, 2026\nDue Date: August 31, 2026"),
        ("Payment Terms", "Amount due within thirty (30) days of the invoice date. Remit payment to North Valley Hospital, Attn: Billing Department. Late payments accrue interest at 1.5% per month."),
    ]
    table = (["Description", "Quantity", "Unit Price", "Amount"], [["Outpatient Consultation", "12", "$185.00", "$2,220.00"], ["Laboratory Services", "8", "$95.00", "$760.00"], ["Imaging - MRI", "2", "$1,240.00", "$2,480.00"]])
    write_pdf(OUT_DIR / fn, "Hospital Invoice", sections, table=table)
    _label(fn, "financial.invoices.hospital_invoice", "financial", "invoice", "hospital_invoice")


def build_reimbursement_policy() -> None:
    fn = "reimbursement_policy_outpatient.pdf"
    sections = [
        ("Purpose", "This reimbursement policy defines how outpatient services are reimbursed under network agreements with Meridian Health Partners and Summit Care Network."),
        ("Reimbursement Rates", "Reimbursement rates for outpatient evaluation and management services are set at 80% of the current fee schedule unless otherwise negotiated in a Provider Agreement."),
        ("Fee Schedule", "The fee schedule is reviewed annually and reimbursement rate changes take effect on January 1 of each calendar year. Providers are notified sixty (60) days in advance of any reimbursement methodology change."),
        ("Appeals Process", "Providers may appeal a reimbursement determination within ninety (90) days of the remittance date by submitting a written appeal to the Reimbursement Review Committee."),
    ]
    write_pdf(OUT_DIR / fn, "Reimbursement Policy: Outpatient Services", sections)
    _label(fn, "financial.reimbursement.reimbursement_policy", "financial", "policy", "reimbursement_policy")


def build_api_architecture_document() -> None:
    fn = "api_architecture_document.md"
    content = """# API Architecture Document: Patient Portal Platform

## Overview

This document describes the system architecture for the Patient Portal API platform, including service boundaries, data flow, and integration points.

## System Architecture

The platform follows a service-oriented architecture with the following components:

- API Gateway (request routing, rate limiting)
- Patient Identity Service
- Appointment Scheduling Service
- Document Retrieval Service
- Notification Service

## API Design

All services expose a REST API secured via OAuth2 bearer tokens. The API design follows resource-oriented URL conventions.

```
GET /api/v1/patients/{patient_id}/appointments
POST /api/v1/appointments
```

## Authentication

Requests must include an `Authorization: Bearer <token>` header. Tokens are issued by the Identity Service and expire after sixty (60) minutes.

## Data Flow

Client requests enter through the API Gateway, are authenticated against the Identity Service, and are routed to the appropriate downstream service. Responses are cached at the gateway layer for fifteen (15) seconds where safe.
"""
    write_text(OUT_DIR / fn, content)
    _label(fn, "technical.architecture.api_architecture_document", "technical", "architecture_document", "api_architecture_document")


def build_production_runbook() -> None:
    fn = "production_runbook_patient_portal.md"
    content = """# Production Runbook: Patient Portal API

## Service Overview

This runbook covers on-call procedures for the Patient Portal API platform.

## On-call Escalation

- Primary on-call: Platform Engineering (PagerDuty rotation)
- Secondary escalation: Engineering Manager after 15 minutes unacknowledged
- Executive escalation: VP Engineering after 60 minutes for Sev1 incidents

## Incident Response Steps

1. Acknowledge the page and open the incident channel.
2. Check the API Gateway dashboard for elevated error rates.
3. Check the Identity Service health endpoint.
4. If the Identity Service is degraded, follow the Rollback Procedure below.

## Rollback Procedure

Roll back to the previous stable deployment using the deployment pipeline's "Rollback" action. Verify health checks pass before closing the incident.

## Troubleshooting

- 401 error spikes: check for expired signing keys in the Identity Service.
- 504 timeouts: check downstream Document Retrieval Service latency.
- Known issue: cache invalidation lag can cause stale appointment data for up to 15 seconds.
"""
    write_text(OUT_DIR / fn, content)
    _label(fn, "technical.runbooks.production_runbook", "technical", "runbook", "production_runbook")


def build_hipaa_compliance_policy() -> None:
    fn = "hipaa_compliance_policy.pdf"
    sections = [
        ("Purpose", "This HIPAA Compliance Policy establishes requirements for protecting Protected Health Information (PHI) in accordance with the HIPAA Privacy Rule and Security Rule."),
        ("Scope", "This compliance policy applies to all workforce members with access to PHI, including clinical, administrative, and technical staff."),
        ("Protected Health Information", "PHI includes any individually identifiable health information transmitted or maintained in any form. Workforce members must access PHI only as required for their job function."),
        ("Security Rule Safeguards", "The organization maintains administrative, physical, and technical safeguards including access controls, audit logging, and encryption of PHI at rest and in transit."),
        ("Breach Notification", "Any suspected breach of PHI must be reported to the Privacy Officer within twenty-four (24) hours of discovery, consistent with the Breach Notification Rule."),
        ("Enforcement", "Violations of this compliance policy may result in disciplinary action up to and including termination, consistent with applicable regulation."),
    ]
    write_pdf(OUT_DIR / fn, "HIPAA Compliance Policy", sections)
    _label(fn, "legal.compliance.hipaa_compliance_policy", "legal", "compliance_policy", "hipaa_compliance_policy")


def build_provider_network_agreement() -> None:
    fn = "provider_network_agreement.docx"
    sections = [
        ("Network Participation", "This Provider Network Agreement governs Provider's in-network participation status with Meridian Health Partners' commercial and Medicare Advantage panels."),
        ("Credentialing Requirements", "Provider must complete initial credentialing and re-credentialing every thirty-six (36) months, including verification of licensure, board certification, and malpractice coverage."),
        ("Panel Status", "Provider panel status may be designated open or closed based on network adequacy needs determined by the Provider Management team."),
        ("Term and Renewal", "This Agreement renews automatically for successive one-year terms unless either party provides ninety (90) days written notice of non-renewal."),
    ]
    write_docx(OUT_DIR / fn, "Provider Network Agreement", sections)
    _label(fn, "legal.agreements.provider_network_agreement", "legal", "agreement", "provider_network_agreement")


def build_clinical_procedure() -> None:
    fn = "clinical_procedure_central_line.html"
    sections: list[tuple[str, int, str | list[str]]] = [
        ("Purpose", 2, "This clinical procedure describes the standard steps for central venous catheter insertion and post-procedure monitoring."),
        ("Indications", 2, ["Long-term IV access", "Hemodynamic monitoring", "Administration of vesicant medications"]),
        ("Pre-Procedure Checklist", 2, "Confirm informed consent, verify site with ultrasound, and complete the central line insertion checklist per protocol."),
        ("Procedure Steps", 2, ["Prepare sterile field and maximal barrier precautions", "Identify vessel via ultrasound guidance", "Insert catheter using Seldinger technique", "Confirm placement via chest x-ray"]),
        ("Post-Procedure Monitoring", 2, "Monitor insertion site for signs of infection every shift. Document dressing changes per protocol."),
    ]
    table = (["Complication", "Rate"], [["Infection", "1.2%"], ["Pneumothorax", "0.8%"], ["Malposition", "2.1%"]])
    write_html(OUT_DIR / fn, "Clinical Procedure: Central Venous Catheter Insertion", sections, table=table)
    _label(fn, "clinical.procedures.clinical_procedure", "clinical", "procedure", "clinical_procedure")


def build_financial_report() -> None:
    fn = "financial_report_q3_2026.xlsx"
    headers = ["Department", "Q1 Revenue", "Q2 Revenue", "Q3 Revenue", "Q3 Payment Volume"]
    rows = [
        ["Outpatient Services", 4200000, 4350000, 4480000, 18650],
        ["Inpatient Services", 8100000, 8250000, 8390000, 3210],
        ["Diagnostic Imaging", 1950000, 2010000, 2085000, 6420],
        ["Emergency Services", 3100000, 3225000, 3310000, 9870],
    ]
    write_xlsx(OUT_DIR / fn, {"Revenue Summary": (headers, rows)})
    _label(fn, "financial.payments.financial_report", "financial", "financial_report", "financial_report")


def build_technical_architecture_presentation() -> None:
    fn = "technical_architecture_presentation.pptx"
    slides = [
        ("Technical Architecture Overview", ["Patient Portal Platform Modernization", "Platform Engineering, 2026"]),
        ("Current State Architecture", ["Monolithic patient portal application", "Single relational database", "Manual deployment process"]),
        ("Target State Architecture", ["Service-oriented API architecture", "Independently scalable services", "Automated CI/CD pipeline"]),
        ("Migration Roadmap", ["Phase 1: Extract Identity Service", "Phase 2: Extract Scheduling Service", "Phase 3: Decommission legacy monolith"]),
        ("Risks", ["Data migration downtime window", "Third-party integration compatibility", "Team ramp-up on new architecture"]),
    ]
    write_pptx(OUT_DIR / fn, slides)
    _label(fn, "technical.architecture.technical_architecture_presentation", "technical", "presentation", "technical_architecture_presentation")


def build_meeting_notes() -> None:
    fn = "meeting_notes_provider_relations.txt"
    content = """MEETING NOTES: PROVIDER RELATIONS SYNC

Date: August 14, 2026

Attendees:
Jordan Ellis (Provider Relations), Priya Natarajan (Network Management), Sam Whitfield (Finance)

Agenda:
Review Q3 provider network changes and outstanding credentialing items.

Discussion:
North Valley Hospital re-credentialing is on track for September completion. Summit Care Network requested an update to their reimbursement rate schedule, to be reviewed by Finance. Two new provider agreements are pending legal review.

Action Items:
Priya to follow up with Summit Care Network on rate schedule by August 21.
Sam to confirm Q3 reimbursement rate changes with Finance leadership.
Jordan to schedule credentialing committee review for North Valley Hospital.

Next Steps:
Next sync scheduled for August 28, 2026.
"""
    write_text(OUT_DIR / fn, content)
    _label(fn, "administrative.communications.meeting_notes", "administrative", "communication", "meeting_notes")


def build_regulatory_document() -> None:
    fn = "regulatory_document_cms_updates.html"
    sections: list[tuple[str, int, str | list[str]]] = [
        ("Overview", 2, "This regulatory document summarizes updated CMS requirements affecting provider reimbursement and compliance reporting for the upcoming plan year."),
        ("Effective Date", 2, "The updated regulation and associated compliance requirements take effect January 1, 2027."),
        ("Key Requirements", 2, ["Updated CPT code reporting requirements", "Revised reimbursement rate disclosure obligations", "New provider directory accuracy attestation"]),
        ("Compliance Deadlines", 2, "Organizations must complete the provider directory attestation by November 1, 2026, in accordance with the regulatory guidance."),
        ("Applicability", 2, "This regulation applies to all providers and payers participating in Medicare Advantage and Medicaid managed care programs."),
    ]
    write_html(OUT_DIR / fn, "Regulatory Update: CMS Reimbursement and Compliance Requirements", sections)
    _label(fn, "legal.regulations.regulatory_document", "legal", "regulatory_document", "regulatory_document")


def build_ambiguous_document() -> None:
    """Deliberately weak-signal document for the low-confidence / human
    review demo scenario (spec section 48). Mixes invoice and reimbursement
    language without a clear heading or dominant signal, so the classifier
    is expected to land in the 0.50-0.65 confidence band."""
    fn = "ambiguous_mixed_memo.pdf"
    sections = [
        ("Memo", "Regarding recent account activity between North Valley Hospital and Summit Care Network for the period ending July 2026."),
        ("", "Amount reflected reconciles against the rate applied for outpatient services during the period. The balance shown may be adjusted pending finance review. Please confirm whether the figure represents the amount due or the amount already applied against the account."),
        ("", "Further details will follow once finance confirms treatment of the balance. No further action is required at this time."),
    ]
    write_pdf(OUT_DIR / fn, "Account Memo", sections)
    _label(fn, "financial.reimbursement.reimbursement_policy", "financial", "unclear", "ambiguous_low_confidence")


def build_claims_export() -> None:
    fn = "claims_export_q3.csv"
    headers = ["claim_line_id", "claim_id", "cpt_code", "denial_code", "billed_amount", "paid_amount"]
    rows = [
        ["L-9001", "CLM-100236", "80053", "CO-16", 95.00, 0.00],
        ["L-9002", "CLM-100239", "99213", "", 185.00, 148.00],
        ["L-9003", "CLM-100240", "36415", "", 25.00, 20.00],
        ["L-9004", "CLM-100241", "99396", "CO-45", 210.00, 175.00],
    ]
    write_csv(OUT_DIR / fn, headers, rows)
    _label(fn, "financial.claims.claims_data_export", "financial", "claims_data_export", "claims_data_export")


def build_provider_outreach_email() -> None:
    fn = "provider_outreach_email.eml"
    body = """Dear Network Provider,

We are writing to inform you of an important network update. Meridian Health Partners is launching a new credentialing portal effective September 1, 2026.

All in-network providers, including North Valley Hospital and affiliated practices, must complete re-attestation through the new portal within sixty (60) days of launch to maintain active network participation status.

Please direct questions to the Provider Relations team.

Thank you,
Provider Network Management
"""
    write_eml(
        OUT_DIR / fn,
        sender="Provider Network Management <network@meridianhealthpartners.example>",
        to="providers@northvalleyhospital.example",
        subject="Network Update: New Credentialing Portal",
        date="Mon, 17 Aug 2026 09:00:00 -0500",
        body=body,
    )
    _label(fn, "administrative.communications.provider_outreach_email", "administrative", "communication", "provider_outreach_email")


BUILDERS = [
    build_provider_agreement,
    build_employee_healthcare_policy,
    build_patient_care_guideline,
    build_insurance_claim_report,
    build_hospital_invoice,
    build_reimbursement_policy,
    build_api_architecture_document,
    build_production_runbook,
    build_hipaa_compliance_policy,
    build_provider_network_agreement,
    build_clinical_procedure,
    build_financial_report,
    build_technical_architecture_presentation,
    build_meeting_notes,
    build_regulatory_document,
    build_ambiguous_document,
    build_claims_export,
    build_provider_outreach_email,
]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for builder in BUILDERS:
        builder()

    labels_path = OUT_DIR / "golden_labels.json"
    labels_path.write_text(json.dumps(GOLDEN_LABELS, indent=2), encoding="utf-8")

    print(f"Generated {len(BUILDERS)} synthetic documents in {OUT_DIR}")
    print(f"Golden labels written to {labels_path}")


if __name__ == "__main__":
    main()
